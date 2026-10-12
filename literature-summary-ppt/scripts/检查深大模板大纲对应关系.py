from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


编号格式 = re.compile(r"^0([1-4])$")


def 遍历文字(形状):
    if 形状.shape_type == MSO_SHAPE_TYPE.GROUP:
        for 子形状 in 形状.shapes:
            yield from 遍历文字(子形状)
        return
    文字 = getattr(形状, "text", "").strip()
    if 文字:
        yield 文字


def 页面文字(页面):
    结果 = []
    for 形状 in 页面.shapes:
        结果.extend(遍历文字(形状))
    return 结果


def 取得中文标题(文字列表, 编号):
    忽略 = {
        编号,
        "Contents.",
        "结构大纲",
        "美美与共 知行合一",
        "Minzu University of China",
    }
    候选 = [
        文字
        for 文字 in 文字列表
        if 文字 not in 忽略
        and re.search(r"[\u4e00-\u9fff]", 文字)
        and not 文字.startswith("深圳大学")
    ]
    return 候选[0] if 候选 else None


def 解析大纲(页面列表):
    for 页面 in 页面列表:
        分组 = []
        for 形状 in 页面.shapes:
            if 形状.shape_type != MSO_SHAPE_TYPE.GROUP:
                continue
            文字列表 = list(遍历文字(形状))
            编号 = [文字 for 文字 in 文字列表 if 编号格式.fullmatch(文字)]
            if len(编号) == 1:
                分组.append((编号[0], 文字列表))
        if len(分组) >= 4:
            return {编号: 取得中文标题(文字列表, 编号) for 编号, 文字列表 in 分组}
    return None


def 解析章节页(页面列表):
    章节页 = {}
    for 页码, 页面 in enumerate(页面列表, start=1):
        文字列表 = 页面文字(页面)
        编号 = [文字 for 文字 in 文字列表 if 编号格式.fullmatch(文字)]
        if len(编号) != 1:
            continue
        当前编号 = 编号[0]
        标题 = 取得中文标题(文字列表, 当前编号)
        if 标题:
            章节页[当前编号] = (标题, 页码)
    return 章节页


def main() -> int:
    解析器 = argparse.ArgumentParser(description="检查深大模板封面以及大纲与章节页标题")
    解析器.add_argument("演示文稿", type=Path)
    解析器.add_argument("--跳过封面检查", action="store_true")
    参数 = 解析器.parse_args()

    演示文稿 = Presentation(参数.演示文稿)
    页面列表 = list(演示文稿.slides)
    错误 = []

    if not 参数.跳过封面检查:
        封面文字 = 页面文字(页面列表[0]) if 页面列表 else []
        if "文献分享" not in 封面文字:
            错误.append("封面主标题不是“文献分享”")

    大纲 = 解析大纲(页面列表)
    if not 大纲:
        错误.append("未找到包含 01 至 04 的四项结构大纲页")
    else:
        章节页 = 解析章节页(页面列表)
        for 编号 in ("01", "02", "03", "04"):
            大纲标题 = 大纲.get(编号)
            对应章节 = 章节页.get(编号)
            if not 大纲标题:
                错误.append(f"大纲缺少 {编号} 的中文标题")
            elif not 对应章节:
                错误.append(f"未找到 {编号} 对应的章节页")
            elif 对应章节[0] != 大纲标题:
                错误.append(
                    f"{编号} 不一致：大纲为“{大纲标题}”，"
                    f"第 {对应章节[1]} 页章节标题为“{对应章节[0]}”"
                )

    if 错误:
        for 信息 in 错误:
            print(f"错误：{信息}")
        return 1

    print("通过：封面与 01 至 04 的大纲和章节标题一致")
    return 0


if __name__ == "__main__":
    sys.exit(main())
